# -*- coding: utf-8 -*-
r"""数据读取器：把三个项目的成品/源读成「可检索的文本行」，供首页助手引用。

设计要点：
1. **零硬依赖**：xlsx / docx 都能用标准库（zipfile + re）解析；装了 openpyxl 就优先用它
   （日期、公式结果更准），没装也不影响——沙箱解释器的依赖经常轮换，不能押在单一库上。
2. **按 mtime 缓存**：成品表动辄上万行，不能每次提问都全量解析一遍。
3. **只带命中行**：提问时按关键词打分，返回最相关的若干行 + 来源路径，不整表灌进模型。
"""
import html
import json
import os
import re
import time
import zipfile

from . import config

# ---------------------------------------------------------------- 解析缓存
_CACHE = {}
_CACHE_MAX = 40


def _sig(path):
    try:
        st = os.stat(path)
        return (st.st_mtime, st.st_size)
    except OSError:
        return None


def _cached(key, path, loader):
    """按 (路径, mtime, size) 缓存解析结果。文件没变就直接复用。"""
    c = config.load().get("datareader") or {}
    if not c.get("index_enabled", True):
        return loader(path)
    sig = _sig(path)
    hit = _CACHE.get(key)
    if hit and hit[0] == sig:
        return hit[1]
    payload = loader(path)
    if len(_CACHE) > _CACHE_MAX:
        _CACHE.pop(next(iter(_CACHE)), None)
    _CACHE[key] = (sig, payload)
    return payload


def _have(mod):
    try:
        __import__(mod)
        return True
    except Exception:
        return False


# ---------------------------------------------------------------- xlsx
_ROW_RE = re.compile(r"<row\b([^>]*?)(?:/>|>(.*?)</row>)", re.S)
_CELL_RE = re.compile(r"<c\b([^>]*?)(?:/>|>(.*?)</c>)", re.S)
_ATTR_RE = re.compile(r'(\w+)="([^"]*)"')
_V_RE = re.compile(r"<v>(.*?)</v>", re.S)
_T_RE = re.compile(r"<t\b[^>]*>(.*?)</t>", re.S)
_SI_RE = re.compile(r"<si>(.*?)</si>|<si\b[^>]*>(.*?)</si>", re.S)
_TAG_RE = re.compile(r"<[^>]+>")


def _col_idx(ref):
    """A1 -> 0，AB12 -> 27。"""
    n = 0
    for ch in (ref or ""):
        if ch.isalpha():
            n = n * 26 + (ord(ch.upper()) - 64)
        else:
            break
    return n - 1


def _shared_strings(zf):
    try:
        xml = zf.read("xl/sharedStrings.xml").decode("utf-8", "replace")
    except KeyError:
        return []
    out = []
    for a, b in _SI_RE.findall(xml):
        out.append(html.unescape(_TAG_RE.sub("", a or b)))
    return out


def _sheet_names(zf):
    try:
        xml = zf.read("xl/workbook.xml").decode("utf-8", "replace")
    except KeyError:
        return []
    return [html.unescape(n) for n in re.findall(r'<sheet\b[^>]*name="([^"]*)"', xml)]


def _parse_sheet(xml, shared, max_rows):
    rows = []
    for _, rbody in _ROW_RE.findall(xml):
        if len(rows) >= max_rows:
            break
        if not rbody:
            continue
        cells = {}
        for cattr, cbody in _CELL_RE.findall(rbody):
            a = dict(_ATTR_RE.findall(cattr))
            idx = _col_idx(a.get("r", "")) if a.get("r") else len(cells)
            t = a.get("t", "n")
            if not cbody:
                continue
            if t == "s":
                m = _V_RE.search(cbody)
                i = int(m.group(1)) if (m and m.group(1).strip().lstrip("-").isdigit()) else -1
                val = shared[i] if 0 <= i < len(shared) else ""
            elif t == "inlineStr":
                val = html.unescape(_TAG_RE.sub("", "".join(_T_RE.findall(cbody))))
            else:
                m = _V_RE.search(cbody)
                val = html.unescape(m.group(1)).strip() if m else ""
            if val != "":
                cells[idx] = val
        if cells:
            w = max(cells) + 1
            rows.append([cells.get(i, "") for i in range(w)])
    return rows


def _read_xlsx_stdlib(path, sheet=None, max_rows=2000):
    with zipfile.ZipFile(path) as zf:
        names = _sheet_names(zf)
        shared = _shared_strings(zf)
        target = "xl/worksheets/sheet1.xml"
        if sheet:
            cands = [n for n in zf.namelist() if n.startswith("xl/worksheets/")]
            if sheet.isdigit():
                idx = int(sheet)
                target = "xl/worksheets/sheet%d.xml" % idx if \
                    ("xl/worksheets/sheet%d.xml" % idx) in cands else (cands[0] if cands else target)
            else:
                i = names.index(sheet) if sheet in names else 0
                target = "xl/worksheets/sheet%d.xml" % (i + 1)
                if target not in cands and cands:
                    target = cands[0]
        try:
            xml = zf.read(target).decode("utf-8", "replace")
        except KeyError:
            cand = [n for n in zf.namelist() if n.startswith("xl/worksheets/") and n.endswith(".xml")]
            if not cand:
                return {"headers": [], "rows": [], "sheets": names, "sheet": ""}
            xml = zf.read(cand[0]).decode("utf-8", "replace")
            target = cand[0]
        rows = _parse_sheet(xml, shared, max_rows)
    headers = rows[0] if rows else []
    body = rows[1:] if rows else []
    return {"headers": headers, "rows": body, "sheets": names, "sheet": sheet or (names[0] if names else "")}


def _read_xlsx_openpyxl(path, sheet=None, max_rows=2000):
    import openpyxl  # noqa
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    names = list(wb.sheetnames)
    ws = wb[sheet] if (sheet and sheet in names) else wb[names[0]]
    rows = []
    for r in ws.iter_rows(values_only=True):
        if len(rows) >= max_rows:
            break
        rows.append(["" if c is None else (str(c) if not isinstance(c, str) else c) for c in r])
    try:
        wb.close()
    except Exception:
        pass
    return {"headers": rows[0] if rows else [], "rows": rows[1:] if rows else [],
            "sheets": names, "sheet": ws.title}


def read_xlsx(path, sheet=None, max_rows=None):
    """读表格，返回 {headers, rows, sheets, sheet, path, truncated, error}。"""
    c = config.load().get("datareader") or {}
    max_rows = max_rows or int(c.get("max_rows", 3000) or 3000)
    path = os.path.abspath(path)
    out = {"path": path.replace("\\", "/"), "headers": [], "rows": [],
           "sheets": [], "sheet": "", "truncated": False, "error": ""}
    if not os.path.isfile(path):
        out["error"] = "文件不存在"
        return out

    def loader(_p):
        res = None
        if _have("openpyxl"):
            try:
                res = _read_xlsx_openpyxl(_p, sheet, max_rows + 1)
            except Exception:
                res = None
        if res is None:
            try:
                res = _read_xlsx_stdlib(_p, sheet, max_rows + 1)
            except Exception as e:
                return {"error": "解析失败：%s" % e}
        return res

    try:
        res = _cached(("xlsx", path, sheet or "", max_rows), path, loader)
    except Exception as e:
        out["error"] = str(e)
        return out
    if res.get("error"):
        out["error"] = res["error"]
        return out
    rows = res.get("rows") or []
    if len(rows) > max_rows:
        rows = rows[:max_rows]
        out["truncated"] = True
    out.update({"headers": res.get("headers") or [], "rows": rows,
                "sheets": res.get("sheets") or [], "sheet": res.get("sheet") or ""})
    return out


# ---------------------------------------------------------------- docx
_P_RE = re.compile(r"<w:p\b[^>]*>(.*?)</w:p>", re.S)
# docx 正文标签带命名空间 <w:t>；通配 (?:w:)?t 同时兼容无前缀写法
_WT_RE = re.compile(r"<(?:w:)?t\b[^>]*>(.*?)</(?:w:)?t>", re.S)


def read_docx(path, max_chars=None):
    """读 Word，返回 {path, title, paragraphs, chars, error}。纯标准库解析。"""
    c = config.load().get("datareader") or {}
    max_chars = max_chars or int(c.get("snippet_chars", 1200) or 1200) * 20
    path = os.path.abspath(path)
    out = {"path": path.replace("\\", "/"), "title": "", "paragraphs": [],
           "chars": 0, "error": ""}
    if not os.path.isfile(path):
        out["error"] = "文件不存在"
        return out

    def loader(_p):
        try:
            with zipfile.ZipFile(_p) as zf:
                xml = zf.read("word/document.xml").decode("utf-8", "replace")
        except Exception as e:
            return {"error": "解析失败：%s" % e}
        paras = []
        for blk in _P_RE.findall(xml):
            txt = html.unescape(_TAG_RE.sub("", "".join(_WT_RE.findall(blk)))).strip()
            if txt:
                paras.append(txt)
        return {"paragraphs": paras}

    res = _cached(("docx", path), path, loader)
    if res.get("error"):
        out["error"] = res["error"]
        return out
    paras = res.get("paragraphs") or []
    out["paragraphs"] = paras
    out["title"] = paras[0][:80] if paras else os.path.basename(path)
    out["chars"] = sum(len(p) for p in paras)
    if out["chars"] > max_chars:   # 截断：只留前若干段
        acc, keep = 0, []
        for p in paras:
            acc += len(p)
            keep.append(p)
            if acc > max_chars:
                break
        out["paragraphs"] = keep
    return out


def read_text(path, max_chars=None):
    c = config.load().get("datareader") or {}
    max_chars = max_chars or int(c.get("snippet_chars", 1200) or 1200) * 8
    path = os.path.abspath(path)
    out = {"path": path.replace("\\", "/"), "text": "", "chars": 0, "error": ""}
    try:
        with open(path, encoding="utf-8", errors="replace") as fh:
            out["text"] = fh.read(max_chars)
        out["chars"] = len(out["text"])
    except Exception as e:
        out["error"] = str(e)
    return out


# ---------------------------------------------------------------- 文件定位
def _dir(key):
    return config.workspace_root(key)


def _list_ext(root, exts, limit=40, recursive=False):
    """列出目录下匹配后缀的文件，新的在前。"""
    root = root or ""
    if not root or not os.path.isdir(root):
        return []
    hits = []
    try:
        if recursive:
            for dp, dns, fns in os.walk(root):
                dns[:] = [d for d in dns if not d.startswith((".", "_", "~"))]
                for f in fns:
                    if os.path.splitext(f)[1].lower() in exts and not f.startswith("~$"):
                        p = os.path.join(dp, f)
                        hits.append((os.path.getmtime(p), p))
        else:
            for f in os.listdir(root):
                p = os.path.join(root, f)
                if os.path.isfile(p) and os.path.splitext(f)[1].lower() in exts \
                        and not f.startswith("~$"):
                    hits.append((os.path.getmtime(p), p))
    except Exception:
        return []
    hits.sort(reverse=True)
    return [p for _, p in hits[:limit]]


def anno_tables(limit=6):
    """公告目录里的整理表（成品/产出），「公告整理」优先，新的在前。"""
    d = _dir("anno")
    xs = _list_ext(d, {".xlsx"}, limit=30)
    xs.sort(key=lambda p: (0 if "公告整理" in os.path.basename(p) else 1,
                           -os.path.getmtime(p)))
    return xs[:limit]


def anno_input_rows():
    """Anno/input.json 里的精读结果（结构化 9 字段），比成品表更快更准。"""
    p = os.path.join(_dir("anno") or "", "input.json")
    if not os.path.isfile(p):
        return []
    try:
        with open(p, encoding="utf-8") as fh:
            data = json.load(fh)
        rows = (data or {}).get("rows") or []
        return [r for r in rows if isinstance(r, dict)]
    except Exception:
        return []


def anno_dumps(limit=50):
    return _list_ext(os.path.join(_dir("anno") or "", "_txt_dump"), {".txt"}, limit=limit)


def data_tables(limit=6):
    """Data 下的总表：根目录 + 各日期目录，新的在前。"""
    d = _dir("data")
    xs = _list_ext(d, {".xlsx"}, limit=40, recursive=True)
    xs.sort(key=lambda p: (0 if "可交换债数据" in os.path.basename(p) else 1,
                           -os.path.getmtime(p)))
    return xs[:limit]


def wechat_docs(limit=20):
    return _list_ext(_dir("wechat"), {".docx"}, limit=limit)


def prompt_file():
    return os.path.join(config.PLATFORM_ROOT, "tools", "eb-anno-digest", "prompt.md")


# ---------------------------------------------------------------- 检索
def _tokens(q):
    """把问题切成可匹配片段：英文/数字串 + 中文整串 + 中文 2-gram。"""
    q = (q or "").strip()
    toks = []
    for m in re.findall(r"[A-Za-z0-9._]{2,}", q):
        toks.append(m.lower())
    cn = re.sub(r"[^\u4e00-\u9fa5]+", "", q)
    if cn:
        toks.append(cn)
        if len(cn) <= 16:
            toks += [cn[i:i + 2] for i in range(len(cn) - 1)]
    seen, out = set(), []
    for t in toks:
        if len(t) >= 2 and t not in seen:
            seen.add(t)
            out.append(t)
    return out


def score_text(text, tokens):
    s = (text or "").lower()
    if not s:
        return 0
    return sum(len(t) for t in tokens if t in s)


def search_rows(headers, rows, query, limit=5):
    """在表格行里按关键词打分，返回 [(score, line_no, row)]。"""
    toks = _tokens(query)
    if not toks:
        return []
    scored = []
    for i, row in enumerate(rows, start=2):      # 第 1 行是表头
        joined = " ".join(str(c) for c in row if c not in (None, ""))
        sc = score_text(joined, toks)
        if sc:
            scored.append((sc, i, row))
    scored.sort(key=lambda x: -x[0])
    return scored[:limit]


def fmt_row(headers, row, max_chars=None):
    c = config.load().get("datareader") or {}
    max_chars = max_chars or int(c.get("snippet_chars", 1200) or 1200)
    parts = []
    for i, v in enumerate(row):
        v = str(v if v is not None else "").strip()
        if not v:
            continue
        h = str(headers[i]).strip() if i < len(headers) else ""
        parts.append(("%s：%s" % (h, v)) if h else v)
    text = " | ".join(parts)
    return text[:max_chars] + ("…" if len(text) > max_chars else "")


def search_in_text(path, query, window=400, limit=3):
    """在纯文本里找关键词首次出现处，返回上下文片段。"""
    toks = _tokens(query)
    if not toks or not os.path.isfile(path):
        return []
    try:
        with open(path, encoding="utf-8", errors="replace") as fh:
            text = fh.read(400000)
    except Exception:
        return []
    low = text.lower()
    hits, pos = [], 0
    for t in sorted(toks, key=len, reverse=True)[:6]:
        p = low.find(t, pos)
        if p >= 0:
            hits.append(text[max(0, p - 120): p + window].strip())
            pos = p + len(t)
        if len(hits) >= limit:
            break
    return hits


# ---------------------------------------------------------------- 分源查询
def _snip(source, title, ref, text, score=0, url=None):
    return {"source": source, "title": title, "ref": ref, "text": text,
            "score": score, "url": url or ""}


def anno_lookup(query, limit=5):
    """查公告：优先 input.json（结构化），其次成品表，最后翻转储正文。"""
    snips = []
    rows = anno_input_rows()
    if rows:
        headers = ["公告日期", "上市公司", "可交换债", "证券代码", "标的股票",
                   "股票代码", "换股价格", "公告性质", "内容摘要"]
        scored = []
        for r in rows:
            joined = " ".join(str(r.get(k, "")) for k in headers)
            sc = score_text(joined, _tokens(query))
            if sc:
                scored.append((sc, r))
        scored.sort(key=lambda x: -x[0])
        for sc, r in scored[:limit]:
            snips.append(_snip("anno", "%s %s" % (r.get("可交换债", ""), r.get("公告日期", "")),
                               "Anno/input.json", fmt_row(headers, [r.get(k, "") for k in headers]), sc))
    if len(snips) < limit:
        for p in anno_tables(limit=2):
            t = read_xlsx(p)
            if t.get("error") or not t["rows"]:
                continue
            for sc, line, row in search_rows(t["headers"], t["rows"], query, limit=limit):
                snips.append(_snip("anno", "%s 第%d行" % (os.path.basename(p), line),
                                   t["path"], fmt_row(t["headers"], row), sc))
            if snips:
                break
    if not snips:
        for p in anno_dumps(limit=30):
            hits = search_in_text(p, query, limit=1)
            if hits:
                snips.append(_snip("anno", os.path.basename(p), p.replace("\\", "/"),
                                   hits[0] + "…", 1))
            if len(snips) >= limit:
                break
    snips.sort(key=lambda s: -s.get("score", 0))
    return snips[:limit]


def data_lookup(query, limit=5):
    snips = []
    for p in data_tables(limit=2):
        t = read_xlsx(p)
        if t.get("error") or not t["rows"]:
            continue
        for sc, line, row in search_rows(t["headers"], t["rows"], query, limit=limit):
            snips.append(_snip("data", "%s 第%d行" % (os.path.basename(p), line),
                               t["path"], fmt_row(t["headers"], row), sc))
        if snips:
            break
    return snips[:limit]


def _wechat_meta(path):
    """读爬虫写的侧车 json（含原文链接），没有则回 None。"""
    mp = path + ".meta.json"
    if os.path.isfile(mp):
        try:
            with open(mp, encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return None
    return None


def wechat_lookup(query, limit=3):
    """查公众号热文原文（Wechat/*.docx）。

    稀疏打分（score_text）命中正常返回；当有可匹配词但全部不命中时，
    退回时间倒序返回前 limit 篇（score=0），避免问法绕一点就完全召回不到。
    """
    toks = _tokens(query)
    cands = []
    for p in wechat_docs(limit=12):
        d = read_docx(p, max_chars=20000)
        if d.get("error") or not d["paragraphs"]:
            continue
        meta = _wechat_meta(p)
        url = (meta or {}).get("url") if meta else ""
        head = "\n".join(d["paragraphs"][:40])
        sc = score_text(head, toks)
        cands.append((sc, _snip("wechat", d["title"][:60], d["path"],
                                "\n".join(d["paragraphs"][1:12])[:1200], sc, url=url)))
    # 零命中兜底：有可匹配词但关键词全不中 → 退回时间倒序，避免完全召回不到
    if toks and not any(s > 0 for s, _ in cands):
        return [c for s, c in cands[:limit]]
    picked = [(s, c) for s, c in cands if s or not query]
    picked.sort(key=lambda x: -x[0])
    return [c for s, c in picked[:limit]]


def digest_lookup(query, limit=2):
    """查公众号舆情的压缩记忆（Wechat/_digest/*.txt）。

    这些 txt 是把旧抓取文章用模型提炼出的结构化精华（时段/情绪/主线/风险），
    供「某时段市场对某主题的看法」这类跨时间问题检索调用。命中全文（精华本就精炼）
    作为片段返回，使压缩记忆可被 RAG 直接引用。文件名（含日期/时段）也参与匹配，
    让「9月/午盘/复盘」等时间-时段词能命中。
    """
    root = _dir("wechat")
    if not root:
        return []
    ddir = os.path.join(root, "_digest")
    if not os.path.isdir(ddir):
        return []
    toks = _tokens(query)
    cands = []
    for name in sorted(os.listdir(ddir), reverse=True):   # 新的时段在前
        if not name.lower().endswith(".txt"):
            continue
        p = os.path.join(ddir, name)
        t = read_text(p, max_chars=4000)
        body = t.get("text") or ""
        if not body:
            continue
        match_text = name[:-4] + "\n" + body   # 文件名也参与打分
        sc = score_text(match_text, toks) if toks else 1
        cands.append((sc, _snip("digest", name[:-4], p.replace("\\", "/"),
                               body[:1500], sc, url="")))
    # 零命中兜底：有可匹配词但关键词全不中 → 退回时间倒序，避免完全召回不到
    if toks and not any(s > 0 for s, _ in cands):
        return [c for s, c in cands[:limit]]
    picked = [(s, c) for s, c in cands if s or not query]
    picked.sort(key=lambda x: -x[0])
    return [c for s, c in picked[:limit]]


# ---------------------------------------------------------------- 晨会 Report
# 数据源：《晨会分享分析总表.xlsx》「标的明细」sheet（日期/标的/细分板块/方向/
# 核心逻辑/预测锚点/观点原话）。这是晨会观点的唯一结构化出口，回测 HTML 也读它。
REPORT_SHEET = "标的明细"
REPORT_HEADERS = ["日期", "标的", "细分板块", "方向", "核心逻辑", "预测锚点", "观点原话"]
# 方向（即评级）→ 多/空/中分组，与回测报告口径一致
BULL_SET = ("强推", "推荐", "关注")
BEAR_SET = ("负面", "提示风险", "谨慎")
NEU_SET = ("中性", "仅提及")


def report_tables(limit=2):
    """Report/_analysis 下的总表，「晨会分享分析总表」优先。"""
    d = _dir("report")
    if not d:
        return []
    xs = _list_ext(os.path.join(d, "_analysis"), {".xlsx"}, limit=20)
    xs = [p for p in xs if not os.path.basename(p).startswith("~$")]
    xs.sort(key=lambda p: (0 if "晨会分享分析总表" in os.path.basename(p) else 1,
                           -os.path.getmtime(p)))
    return xs[:limit]


def report_rows():
    """「标的明细」全表（带缓存）。返回 {path, headers, rows} 或 None。"""
    xs = report_tables(2)
    if not xs:
        return None
    p = xs[0]

    def loader(_p):
        t = read_xlsx(_p, sheet=REPORT_SHEET, max_rows=5000)
        if t.get("error") or not t.get("rows"):
            return None
        return t

    t = _cached("report_rows", p, loader)
    if not t:
        return None
    return {"path": p.replace("\\", "/"), "headers": t["headers"], "rows": t["rows"]}


def _row_text(row, headers):
    parts = []
    for i, v in enumerate(row):
        v = str(v if v is not None else "").strip()
        if not v:
            continue
        h = headers[i] if i < len(headers) else ""
        parts.append(("%s：%s" % (h, v)) if h else v)
    return " | ".join(parts)


def report_lookup(query, limit=5):
    """晨会观点检索：标的 / 细分板块 / 评级（方向）/ 逻辑 / 观点原话。

    同一标的最多 2 条，避免一只票刷屏；片段里带上日期与评级，便于回答「什么时候说的」。
    """
    t = report_rows()
    if not t:
        return []
    headers = t["headers"] if t["headers"] and any(t["headers"]) else REPORT_HEADERS
    rows = t["rows"]
    scored = search_rows(headers, rows, query, limit=limit * 6)
    if not scored:
        return []
    per_stock, snips = {}, []
    for sc, line, row in scored:
        stock = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        if stock and per_stock.get(stock, 0) >= 2:
            continue
        per_stock[stock] = per_stock.get(stock, 0) + 1
        title = "%s %s %s" % (str(row[0] or ""), stock,
                              str(row[3] or "") if len(row) > 3 else "")
        snips.append(_snip("morning", title.strip(), t["path"],
                           _row_text(row, headers), sc))
        if len(snips) >= limit:
            break
    return snips


def _sector_stats(rows, headers):
    """按「细分板块」聚合：条数 / 多空中分布 / 最新日期 / 代表标的。"""
    idx = {h: i for i, h in enumerate(headers)}
    gi = {k: idx.get(k, -1) for k in ("日期", "标的", "细分板块", "方向")}
    agg = {}
    for r in rows:
        sec = str(r[gi["细分板块"]]).strip() if gi["细分板块"] >= 0 and len(r) > gi["细分板块"] else ""
        if not sec:
            continue
        d = str(r[gi["日期"]]).strip() if gi["日期"] >= 0 and len(r) > gi["日期"] else ""
        st = str(r[gi["标的"]]).strip() if gi["标的"] >= 0 and len(r) > gi["标的"] else ""
        dd = str(r[gi["方向"]]).strip() if gi["方向"] >= 0 and len(r) > gi["方向"] else ""
        a = agg.setdefault(sec, {"n": 0, "bull": 0, "bear": 0, "neu": 0,
                                 "last": "", "stocks": [], "ratings": {}})
        a["n"] += 1
        if dd in BULL_SET:
            a["bull"] += 1
        elif dd in BEAR_SET:
            a["bear"] += 1
        elif dd in NEU_SET:
            a["neu"] += 1
        if dd:
            a["ratings"][dd] = a["ratings"].get(dd, 0) + 1
        if d > a["last"]:
            a["last"] = d
        if st and st not in a["stocks"]:
            a["stocks"].append(st)
    return agg


def _fmt_sector(sec, a, top=8):
    ratings = "、".join("%s×%d" % (k, v) for k, v in
                       sorted(a["ratings"].items(), key=lambda x: -x[1]))
    return ("%s：%d 条（看多 %d / 看空 %d / 中性 %d）｜最新 %s｜评级 %s｜标的 %s"
            % (sec, a["n"], a["bull"], a["bear"], a["neu"], a["last"],
               ratings or "—", "、".join(a["stocks"][:top]) +
               ("…" if len(a["stocks"]) > top else "")))


def report_sector_stats(query="", limit_top=10):
    """板块聚合：命中某板块就给该板块详情，另外总给一份「最被看好」Top 榜。

    回答「XX 板块怎么看」「哪个板块最被看好」「半导体最近什么评级」这类汇总问题。
    """
    t = report_rows()
    if not t:
        return []
    headers = t["headers"] if t["headers"] and any(t["headers"]) else REPORT_HEADERS
    agg = _sector_stats(t["rows"], headers)
    if not agg:
        return []
    snips = []

    # ① 问题里点名的板块
    if query:
        hits = [s for s in agg if s and s in query] or \
               [s for s in agg if s and any(tok == s for tok in _tokens(query))]
        for sec in sorted(set(hits), key=lambda s: -agg[s]["n"])[:3]:
            snips.append(_snip("morning", "板块聚合 · %s" % sec, t["path"],
                               _fmt_sector(sec, agg[sec], top=12), 3))

    # ② 全量榜：按看多类条数排（回答「最看好哪个板块」）
    top = sorted(agg.items(), key=lambda kv: (-kv[1]["bull"], -kv[1]["n"]))[:limit_top]
    if top:
        body = "；\n".join(_fmt_sector(s, a, top=5) for s, a in top)
        total = sum(a["n"] for a in agg.values())
        head = "全库 %d 个细分板块 / %d 条观点。看多类（强推+推荐+关注）最多的板块：\n" % (
            len(agg), total)
        snips.append(_snip("morning", "晨会板块热度 Top%d" % len(top), t["path"],
                           head + body, 2))
    return snips


# ---------------------------------------------------------------- 观点命中回测（HTML 内嵌数据）
# 回测报告 Report/_analysis/F观点命中回测_板块.html 把数据内嵌在 `const DATA = {...}` 里，
# 只按 xlsx 总表检索会漏掉「命中率 / 个股原话」——这两项只有回测报告里有，故单独做一个源。
_BT_CACHE = {"sig": None, "data": None}
_BT_NAME = "F观点命中回测_板块.html"


def _bt_path():
    d = _dir("report")
    return os.path.join(d, "_analysis", _BT_NAME) if d else ""


def _bt_data():
    """解析回测 HTML 内嵌的 `const DATA`，按 (mtime,size) 缓存。

    结构：sectors[] → 板块含 hit_n/hit_d（命中数/总数）、mentions、ret、dir_counts；
    stocks[] → 个股含 code/name/hit_n/hit_d/ret/quotes（观点原话 + 日期 qdate）。
    """
    p = _bt_path()
    if not p or not os.path.isfile(p):
        return None
    sig = _sig(p)
    if _BT_CACHE["data"] is not None and _BT_CACHE["sig"] == sig:
        return _BT_CACHE["data"]
    try:
        with open(p, encoding="utf-8", errors="replace") as fh:
            doc = fh.read()
    except Exception:
        return None
    m = re.search(r"const\s+DATA\s*=\s*", doc)
    if not m:
        return None
    i = m.end()
    while i < len(doc) and doc[i] in " \r\n\t":
        i += 1
    if i >= len(doc) or doc[i] != "{":
        return None
    depth, end = 0, -1
    for j in range(i, len(doc)):
        if doc[j] == "{":
            depth += 1
        elif doc[j] == "}":
            depth -= 1
            if depth == 0:
                end = j
                break
    if end < 0:
        return None
    try:
        data = json.loads(doc[i:end + 1])
    except Exception:
        return None
    _BT_CACHE["sig"], _BT_CACHE["data"] = sig, data
    return data


def _bt_rate(item):
    """命中率：(比率, hit_n, hit_d)。无样本时比率 -1，排序时垫底。"""
    hd = item.get("hit_d") or 0
    hn = item.get("hit_n") or 0
    return (hn / hd if hd else -1.0, hn, hd)


def _bt_sortkey(item):
    """排序键：命中率降序 → 样本数降序 → 区间收益降序。

    小样本（1/1）很容易并列 100%，只看命中率会有一大片并列，
    所以依次用「样本多者更可信」「收益高者更有价值」来打破平局。
    """
    r, _, hd = _bt_rate(item)
    try:
        ret = float(item.get("ret"))
    except (TypeError, ValueError):
        ret = -999.0
    return (-r, -hd, -ret)


def _bt_fmt_sector(s, top=8, quote_top=2, qchars=120):
    """板块片段：命中率 + 评级分布 + 个股命中率排行（前几名附观点原话）。

    片段会被截断到 snippet_chars（默认 1200），所以顺序上先给紧凑排行、再给原话，
    保证「最关键的排名 + 榜首原话」不会被截掉。
    """
    r, hn, hd = _bt_rate(s)
    rate = ("%.0f%%" % (r * 100)) if hd else "—"
    lines = ["%s：命中率 %d/%d = %s｜观点数 %s｜区间收益 %s%%｜评级 %s"
             % (s.get("name") or "-", hn, hd, rate, s.get("mentions") or 0,
                s.get("ret"), json.dumps(s.get("dir_counts") or {}, ensure_ascii=False)),
             "  个股命中率排行（同率按样本数降序）："]
    stocks = sorted(s.get("stocks") or [], key=_bt_sortkey)
    for k, x in enumerate(stocks[:top]):
        xr, xhn, xhd = _bt_rate(x)
        xs = ("%.0f%%" % (xr * 100)) if xhd else "—"
        lines.append("   %d) %s(%s) %d/%d=%s 收益%s%%"
                     % (k + 1, x.get("name") or "-", x.get("code") or "-",
                        xhn, xhd, xs, x.get("ret")))
    for k, x in enumerate(stocks[:quote_top]):
        qs = [q for q in (x.get("quotes") or []) if q][:3]
        if not qs:
            continue
        lines.append("  ▸%s 的观点原话：" % (x.get("name") or "-"))
        for q in qs:
            if isinstance(q, dict):
                lines.append("    [%s] %s" % (q.get("qdate") or "-",
                                              (q.get("quote") or "")[:qchars]))
            else:
                lines.append("    %s" % str(q)[:qchars])
    return "\n".join(lines)


def backtest_lookup(query="", limit_top=10):
    """观点命中回测检索：命中率榜 + 问题点名板块的个股明细（含观点原话）。"""
    data = _bt_data()
    if not data:
        return []
    sectors = data.get("sectors") or []
    if not sectors:
        return []
    q = query or ""
    ref = _bt_path().replace("\\", "/")
    snips = []

    # ① 问题里点名的板块（先整名命中，再退到分词）
    named = [s for s in sectors if s.get("name") and s["name"] in q]
    if not named:
        toks = set(_tokens(q))
        named = [s for s in sectors if s.get("name") and s["name"] in toks]
    for s in named[:3]:
        snips.append(_snip("backtest", "回测 · %s" % s.get("name"), ref,
                           _bt_fmt_sector(s), 4))

    # ② 全板块命中率榜：同率按样本数降序，避免 1/1 的小样本刷榜
    ranked = [s for s in sectors if (s.get("hit_d") or 0) > 0]
    ranked.sort(key=_bt_sortkey)
    if ranked:
        show = ranked[:limit_top]
        body = "；\n".join("%s %d/%d=%.0f%%" % (s.get("name"), s.get("hit_n") or 0,
                                                s.get("hit_d"), _bt_rate(s)[0] * 100)
                          for s in show)
        snips.append(_snip("backtest", "板块命中率 Top%d" % len(show), ref,
                           "全库 %d 个板块，按命中率排序：\n%s" % (len(sectors), body), 2))
    return snips


def overview():
    """工作台总览（给「总览/状态」这类问题用），零 LLM 成本。"""
    anno_dir, data_dir, wechat_dir = _dir("anno"), _dir("data"), _dir("wechat")
    pdfs = _list_ext(anno_dir, {".pdf"}, limit=200)
    dumps = anno_dumps(limit=500)
    rows = anno_input_rows()
    tables = anno_tables(limit=20)
    dtables = data_tables(limit=20)
    docs = wechat_docs(limit=200)
    dates = []
    if data_dir and os.path.isdir(data_dir):
        dates = sorted([d for d in os.listdir(data_dir)
                        if re.match(r"^\d{8}$", d) and
                        os.path.isdir(os.path.join(data_dir, d))], reverse=True)
    return {
        "anno": {"dir": (anno_dir or "").replace("\\", "/"), "pdf_count": len(pdfs),
                 "dump_count": len(dumps), "input_rows": len(rows),
                 "latest_table": (tables[0].replace("\\", "/") if tables else "")},
        "data": {"dir": (data_dir or "").replace("\\", "/"), "date_count": len(dates),
                 "latest_date": dates[0] if dates else "",
                 "latest_table": (dtables[0].replace("\\", "/") if dtables else "")},
        "wechat": {"dir": (wechat_dir or "").replace("\\", "/"), "doc_count": len(docs),
                   "latest": (docs[0].replace("\\", "/") if docs else "")},
        "updated": time.strftime("%Y-%m-%d %H:%M:%S"),
    }
