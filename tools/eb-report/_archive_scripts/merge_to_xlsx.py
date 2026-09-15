# -*- coding: utf-8 -*-
"""把 anno_json/*.json（schema.md 的主记录数组）并入《晨会分享分析总表.xlsx》。

这是管线里原先缺失的一环：LLM 提取产出 anno_json，但没有任何脚本把它们并进总表。

规则
- 主表一行 = 一篇晨会；标的明细每行 = 该篇的一个标的。
- **按日期去重**：总表 主表 里已有的日期默认跳过（幂等，可反复跑）。
- 写盘前自动备份一份 .bak_<时间戳>，不删任何东西（沙箱批量删除守卫：只用 copy/rename）。
- etf明细 / 统计 两张表本脚本不生成（etf 需人工判定，统计是主题分布），保持原样。

用法:
  python merge_to_xlsx.py [--anno <anno_json目录>] [--xlsx <总表.xlsx>] [--force] [--dry]
"""
import os
import re
import sys
import json
import shutil
import argparse
import datetime

try:
    import openpyxl
except ImportError:
    print("需要 openpyxl：请用带依赖的解释器运行（Doubao 沙箱 Python）")
    sys.exit(1)

MAIN_COLS = ['日期', '演讲者', '段落数', '字数', '一级主题', '内容模块', '核心增量信息',
             '提及标的数', '预测条数', '风险点', '外部信源', '质量标记', '标的明细条数']
STOCK_COLS = ['日期', '标的', '细分板块', '方向', '核心逻辑', '预测锚点', '观点原话']


def norm_date(v):
    """归一到 8 位 YYYYMMDD；认不出返回 ''。"""
    if v is None:
        return ''
    s = str(v).strip()
    if re.fullmatch(r'\d{8}', s):
        return s
    if re.fullmatch(r'\d{4}-\d{2}-\d{2}', s):
        return s.replace('-', '')
    if re.fullmatch(r'\d{4}\.\d{1,2}\.\d{1,2}', s):
        y, m, d = re.split(r'[.]', s)
        return '%04d%02d%02d' % (int(y), int(m), int(d))
    return ''


def join_list(v):
    if isinstance(v, (list, tuple)):
        return '、'.join(str(x) for x in v if x)
    return '' if v is None else str(v)


def load_records(anno_dir):
    """读 anno_json 下所有 json，返回 (records, 文件数, 解析失败数)。"""
    recs, nfiles, nbad = [], 0, 0
    if not os.path.isdir(anno_dir):
        return recs, 0, 0
    for name in sorted(os.listdir(anno_dir)):
        if not name.lower().endswith('.json'):
            continue
        path = os.path.join(anno_dir, name)
        nfiles += 1
        try:
            with open(path, encoding='utf-8') as fh:
                data = json.load(fh)
        except Exception as exc:
            print('  !! 读取失败 %s: %s' % (name, exc))
            nbad += 1
            continue
        if isinstance(data, dict):
            data = [data]
        if not isinstance(data, list):
            print('  !! 顶层不是数组，跳过 %s' % name)
            nbad += 1
            continue
        for r in data:
            if isinstance(r, dict):
                recs.append(r)
    return recs, nfiles, nbad


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--anno', default=os.environ.get('REPORT_ANNO_DIR', ''),
                    help='anno_json 目录，默认 <REPORT_DIR>/anno_json')
    ap.add_argument('--xlsx', default='', help='总表 xlsx，默认 <REPORT_DIR>/晨会分享分析总表.xlsx')
    ap.add_argument('--force', action='store_true', help='已存在日期也重新并入（先删旧行）')
    ap.add_argument('--dry', action='store_true', help='只预览不写盘')
    args = ap.parse_args()

    root = os.environ.get('REPORT_DIR', r'C:\Users\<用户名>\Desktop\Report\_analysis')
    anno_dir = args.anno or os.path.join(root, 'anno_json')
    xlsx = args.xlsx or os.path.join(root, '晨会分享分析总表.xlsx')

    if not os.path.exists(xlsx):
        print('找不到总表: %s' % xlsx)
        return 1
    recs, nfiles, nbad = load_records(anno_dir)
    print('anno_json: %s -> %d 个文件, %d 条主记录, %d 个坏文件' % (anno_dir, nfiles, len(recs), nbad))
    if not recs:
        print('没有可并入的记录')
        return 0

    wb = openpyxl.load_workbook(xlsx)
    if '主表' not in wb.sheetnames or '标的明细' not in wb.sheetnames:
        print('总表缺 主表/标的明细 工作表（现有: %s）' % wb.sheetnames)
        return 1
    ws_main = wb['主表']
    ws_stock = wb['标的明细']

    # 已入库日期（主表 A 列）
    have = set()
    for row in ws_main.iter_rows(min_row=2, max_col=1, values_only=True):
        d = norm_date(row[0])
        if d:
            have.add(d)
    print('总表已有日期: %d 个' % len(have))

    if args.force:
        # 删掉待并入日期的旧行（从后往前删，避免行号漂移）
        targets = {norm_date(r.get('date')) for r in recs} - {''}
        for ws, ncol in ((ws_main, 1), (ws_stock, 1)):
            for i in range(ws.max_row, 1, -1):
                if norm_date(ws.cell(i, ncol).value) in targets:
                    ws.delete_rows(i)
        have -= targets
        print('--force: 已清除 %d 个日期的旧行' % len(targets))

    added_main = added_stock = skipped = 0
    for r in recs:
        d = norm_date(r.get('date') or r.get('文件名') or '')
        if not d:
            print('  -- 无日期，跳过一条记录')
            skipped += 1
            continue
        if d in have:
            skipped += 1
            continue
        stocks = r.get('stocks') or []
        forecasts = r.get('forecasts') or []
        ws_main.append([
            d,
            r.get('speaker', '') or '',
            r.get('paras', ''),
            r.get('chars', ''),
            join_list(r.get('themes')),
            join_list(r.get('modules')),
            r.get('core', '') or '',
            r.get('targets_n', len(stocks)),
            len(forecasts),
            join_list(r.get('risks')),
            join_list(r.get('sources')),
            join_list(r.get('quality')),
            len(stocks),
        ])
        added_main += 1
        for s in stocks:
            if not isinstance(s, dict):
                continue
            ws_stock.append([
                d,
                s.get('name', '') or '',
                s.get('sector', '') or '',
                s.get('dir', '') or '',
                s.get('logic', '') or '',
                s.get('anchor', '') or '',
                s.get('quote', '') or '',
            ])
            added_stock += 1
        have.add(d)
        print('  ++ %s: 主表1行 / 标的明细%d行' % (d, len(stocks)))

    print('并入: 主表 %d 行, 标的明细 %d 行, 跳过(已存在) %d 条' % (added_main, added_stock, skipped))
    if args.dry:
        print('--dry: 未写盘')
        return 0
    if added_main == 0:
        print('无新增，不写盘')
        return 0

    bak = xlsx + '.bak_' + datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    shutil.copy2(xlsx, bak)
    print('已备份: %s' % bak)
    wb.save(xlsx)
    print('已写回: %s' % xlsx)
    return 0


if __name__ == '__main__':
    sys.exit(main())
