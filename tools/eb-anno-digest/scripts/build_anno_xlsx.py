# -*- coding: utf-8 -*-
r"""可交换债公告整理表 —— 由 JSON 生成排版好的 xlsx

用法：
    python build_anno_xlsx.py input.json                    # 输出 <input 同名>_公告整理.xlsx
    python build_anno_xlsx.py input.json -o 公告整理.xlsx   # 指定输出
    python build_anno_xlsx.py --check input.json            # 只做校验不产出

输入 JSON 结构（UTF-8）：
{
  "meta": {
    "notes": ["处理说明第 1 行", "处理说明第 2 行"]   // 可选，写入「处理说明」页
  },
  "rows": [
    {
      "公告日期": "2026.8.28",
      "上市公司": "深圳华强",
      "可交换债": "26华强E1",
      "证券代码": "117257.SZ",
      "标的股票": "深圳华强",
      "股票代码": "000062.SZ",
      "换股价格": "31.36元/股",
      "公告性质": "深圳华强集团有限公司...更名公告",
      "内容摘要": "本期债券申报名称为...",     // 130~360 字，事件/要点式
      "条款长摘要": "..."                      // 可选；填了会存到「条款长摘要（原版）」页
    }
  ]
}

也可传入 CSV（表头同上中文列名），自动转 JSON 处理。

输出 3 个 Sheet：
    Sheet1                —— 整理表（9 列，可直接整行粘进成品大表）
    条款长摘要（原版）      —— 仅当任一行填了「条款长摘要」时生成
    处理说明              —— meta.notes 的逐行内容

排版沿用既定偏好：宋体 11 号、不加粗、无填充、无边框、字体纯黑、
不自动换行、冻结首行。列宽对齐成品表，便于整行复制粘贴。
"""
import argparse
import csv
import io
import json
import os
import sys

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border
except ImportError:
    print("=" * 58)
    print("缺少依赖库 openpyxl，请先执行：")
    print("    pip install openpyxl")
    print("=" * 58)
    sys.exit(1)

HEADERS = ["公告日期", "上市公司", "可交换债", "证券代码", "标的股票", "股票代码",
           "换股价格", "公告性质", "内容摘要"]
CENTER_COLS = {1, 2, 3, 4, 5, 6, 7}
WIDTHS = {"A": 11.2, "B": 10.9, "C": 21.9, "D": 20.0, "E": 11.7, "F": 16.5,
          "G": 20.9, "H": 47.3, "I": 149.7}

SUMMARY_MIN, SUMMARY_MAX = 130, 360

FONT = Font(name="宋体", size=11, bold=False, color="000000")
ALIGN_C = Alignment(horizontal="center", vertical="center", wrap_text=False)
ALIGN_L = Alignment(horizontal="left", vertical="center", wrap_text=False)


def load_input(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".json":
        with open(path, encoding="utf-8") as fh:
            data = json.load(fh)
        # 兼容两种写法：{"rows":[...], "meta":{...}} 或直接是行列表 [...]
        if isinstance(data, list):
            return {"meta": {}, "rows": data}
        return data
    if ext in (".csv", ".txt"):
        with open(path, encoding="utf-8-sig", newline="") as fh:
            rows = [dict(r) for r in csv.DictReader(fh)]
        return {"meta": {}, "rows": rows}
    sys.exit(f"[错误] 不支持的输入格式：{ext}（只接受 .json / .csv）")


def validate(rows):
    """返回问题列表。摘要字数只是提醒，不阻断生成。"""
    problems = []
    for i, r in enumerate(rows, start=1):
        for h in HEADERS:
            if not str(r.get(h, "")).strip():
                problems.append(f"第 {i} 行：缺少必填字段「{h}」")
        s = str(r.get("内容摘要", "")).strip()
        n = len(s)
        if n and not (SUMMARY_MIN <= n <= SUMMARY_MAX):
            tag = "偏短" if n < SUMMARY_MIN else "偏长"
            problems.append(f"第 {i} 行（{r.get('可交换债', '?')}）：摘要 {n} 字，{tag}"
                            f"（建议 {SUMMARY_MIN}~{SUMMARY_MAX} 字）")
    seen = {}
    for i, r in enumerate(rows, start=1):
        key = (r.get("可交换债", ""), r.get("公告性质", ""))
        if key in seen:
            problems.append(f"第 {i} 行与第 {seen[key]} 行疑似重复：{key[0]} / {key[1][:30]}")
        else:
            seen[key] = i
    return problems


def style(sheet, ncol, center_cols):
    for r in range(1, sheet.max_row + 1):
        for c in range(1, ncol + 1):
            cell = sheet.cell(row=r, column=c)
            cell.font = FONT
            cell.fill = PatternFill(fill_type=None)
            cell.border = Border()
            cell.alignment = ALIGN_C if c in center_cols else ALIGN_L
        sheet.row_dimensions[r].height = 20


def build(data, out_path):
    rows = data.get("rows", [])
    meta = data.get("meta", {})
    if not rows:
        sys.exit("[错误] 输入中没有任何数据行")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    ws.append(HEADERS)
    for r in rows:
        ws.append([str(r.get(h, "")).strip() for h in HEADERS])
    style(ws, len(HEADERS), CENTER_COLS)
    for k, v in WIDTHS.items():
        ws.column_dimensions[k].width = v
    ws.freeze_panes = "A2"

    # 条款长摘要留存页
    if any(str(r.get("条款长摘要", "")).strip() for r in rows):
        ws2 = wb.create_sheet("条款长摘要（原版）")
        ws2.append(["序号", "可交换债", "证券代码", "公告性质", "条款长摘要"])
        for i, r in enumerate(rows, start=1):
            ws2.append([i, r.get("可交换债", ""), r.get("证券代码", ""),
                        r.get("公告性质", ""), r.get("条款长摘要", "")])
        style(ws2, 5, {1, 2, 3})
        for k, v in {"A": 6, "B": 16, "C": 14, "D": 47.3, "E": 149.7}.items():
            ws2.column_dimensions[k].width = v
        ws2.freeze_panes = "A2"

    # 处理说明
    notes = meta.get("notes", [])
    if notes:
        ws3 = wb.create_sheet("处理说明")
        for n in notes:
            ws3.append([n])
        style(ws3, 1, set())
        ws3.column_dimensions["A"].width = 130
        ws3.freeze_panes = "A2"

    wb.save(out_path)

    lens = [len(str(r.get("内容摘要", "")).strip()) for r in rows]
    print(f"[完成] {out_path}")
    print(f"       数据行 {len(rows)}，摘要字数 min/max/avg = "
          f"{min(lens)}/{max(lens)}/{sum(lens) / len(lens):.0f}")
    print(f"       Sheet：{', '.join(wb.sheetnames)}")


def main():
    ap = argparse.ArgumentParser(description="由 JSON/CSV 生成可交换债公告整理 xlsx")
    ap.add_argument("input", help="输入 .json 或 .csv")
    ap.add_argument("-o", "--out", default=None, help="输出 xlsx 路径")
    ap.add_argument("--check", action="store_true", help="只校验不生成文件")
    ap.add_argument("--strict", action="store_true", help="校验有问题时拒绝生成")
    args = ap.parse_args()

    data = load_input(args.input)
    rows = data.get("rows", [])
    problems = validate(rows)

    if problems:
        print(f"[校验] 发现 {len(problems)} 处问题：")
        for p in problems:
            print(f"       - {p}")
        if args.strict:
            sys.exit("[中止] --strict 模式，请先修正后再生成")
    else:
        print(f"[校验] 通过，共 {len(rows)} 行")

    if args.check:
        return 0

    out = args.out
    if not out:
        stem = os.path.splitext(os.path.basename(args.input))[0]
        out = os.path.join(os.path.dirname(os.path.abspath(args.input)),
                           f"{stem}_公告整理.xlsx")
    build(data, out)
    return 0


if __name__ == "__main__":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
    sys.exit(main())
